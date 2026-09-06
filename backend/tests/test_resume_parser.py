import io

from docx import Document

from app.services.resume_parser import extract_sections, detect_sections, parse_resume_file


def test_extract_sections():
    text = """John Doe
    Software Engineer

    SUMMARY
    Experienced engineer.

    EXPERIENCE
    Acme Corp - Engineer
    - Built things

    SKILLS
    Python, FastAPI
    """
    sections = extract_sections(text)
    assert "summary" in sections
    assert "experience" in sections
    assert "skills" in sections
    assert "Experienced engineer." in sections["summary"]


def test_detect_sections():
    text = "PROFESSIONAL SUMMARY\nExperienced.\nWORK EXPERIENCE\nAcme."
    detected = detect_sections(text)
    assert "Professional summary" in detected
    assert "Work experience" in detected


def test_parse_docx():
    doc = Document()
    doc.add_heading("Jane Doe", 0)
    doc.add_paragraph("Summary: 5 years of experience.")
    buf = io.BytesIO()
    doc.save(buf)

    text, sections, layout = parse_resume_file(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "Jane Doe" in text
    assert layout == {"hasTables": False, "singleColumn": None}